"""
Renders a generate_proposal() sections dict into a formatted .docx file
for email attachment. Must handle PARTIAL section dicts too — the
lightweight WATCH-tier path only produces cover_letter and
executive_summary, not all 10 keys.
"""
import os
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from loguru import logger

from utils.money_scrub import (
    contains_financial_disclosure,
    strip_financial_table_headers,
    strip_monetary_amounts,
)

# Order matters — this is the real proposal structure, not dict
# insertion order. The 4 grouped generator functions map onto these
# slots (org_profile_and_track_record covers what would otherwise be
# two separate PROPOSAL_STRUCTURE entries, etc.) — see proposal_writer.py.
SECTION_ORDER = [
    ("cover_letter", "Cover Letter / Letter of Interest"),
    ("firm_profile", "Presentation of Cortech Consulting Group"),
    ("understanding", "Our Understanding of the Assignment"),
    ("approach_summary", "Proposed Technical Approach — Summary"),
    ("relevant_experience", "Relevant Experience"),
    ("key_experts", "Resources in Staff"),
    ("eligibility", "Eligibility"),
    ("compliance_matrix", "Capability Matrix"),
    ("executive_summary", "Executive Summary"),
    ("org_profile_and_track_record", "Organisational Profile & Track Record"),
    ("introduction_and_framework", "Introduction, Background & Conceptual Framework"),
    ("methodology", "Technical Methodology"),
    ("analysis_plan", "Sampling Strategy & Data Analysis Plan"),
    ("qa_and_ethics", "Quality Assurance & Ethical Considerations"),
    ("risk_register", "Risk Management"),
    ("team_section", "Proposed Team & Core Experts"),
    ("work_plan", "Work Plan & Timeline"),
]


# "- item", "* item", "+ item" — the marker must be followed by whitespace so
# a "---" rule and an "*emphasised*" line opener are not read as bullets.
_BULLET_RE = re.compile(r"^[-*+]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\d+[.)]\s+(.*)$")
_HR_RE = re.compile(r"^([-*_])\1{2,}$")


def _client_safe_section(key: str, content: str) -> str:
    """The Word file is the technical/EOI envelope — it must not carry price."""
    text, _ = strip_monetary_amounts(content)
    text, _ = strip_financial_table_headers(text)
    if contains_financial_disclosure(text):
        raise ValueError(
            f"Refusing to write [{key}] into the technical/EOI Word file: "
            "financial information remains"
        )
    return text


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def _is_table_separator(line: str) -> bool:
    s = line.strip().replace(" ", "")
    return s.startswith("|") and bool(s) and set(s) <= set("|:-")


# Inline emphasis, longest marker first so *** wins over ** wins over *.
# Underscore emphasis is deliberately NOT supported: snake_case tokens that
# legitimately appear in proposals (beneficiary_id, raw_dataset) would be
# silently italicised and lose their underscores.
_INLINE_RE = re.compile(
    r"(\*\*\*[^*]+?\*\*\*"      # ***bold italic***
    r"|\*\*.+?\*\*"             # **bold**
    r"|(?<!\*)\*(?!\s)[^*]+?\*" # *italic*
    r"|`[^`]+`)",               # `code`
    re.S,
)


def _add_formatted_text(paragraph, text: str) -> None:
    """Render inline markdown emphasis into separate Word runs."""
    for token in _INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("***") and token.endswith("***") and len(token) > 6:
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
        elif token.startswith("**") and token.endswith("**") and len(token) > 4:
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            paragraph.add_run(token[1:-1]).font.name = "Consolas"
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


# Kept as the historical name used by _write_table.
_add_text_with_bold = _add_formatted_text


def _add_list_item(doc: Document, text: str, indent: int, numbered: bool) -> None:
    """
    Add a bullet/numbered item at a nesting depth derived from indentation.
    Falls back to a plain paragraph with a literal bullet if the template
    is missing the List styles.
    """
    level = min(indent // 2, 2)
    base = "List Number" if numbered else "List Bullet"
    style = base if level == 0 else f"{base} {level + 1}"
    try:
        paragraph = doc.add_paragraph(style=style)
    except KeyError:
        paragraph = doc.add_paragraph()
        paragraph.add_run("• " if not numbered else "- ")
    _add_formatted_text(paragraph, text)


def _write_table(doc: Document, rows: list[str]) -> None:
    parsed = []
    for row in rows:
        if _is_table_separator(row):
            continue
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        parsed.append(cells)
    if not parsed:
        return
    cols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.style = "Table Grid"
    for ri, row in enumerate(parsed):
        for ci in range(cols):
            cell = table.rows[ri].cells[ci]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _add_text_with_bold(paragraph, row[ci] if ci < len(row) else "")
            if ri == 0:
                for run in paragraph.runs:
                    run.bold = True


def _write_markdown_content(doc: Document, content: str) -> None:
    """
    Render Claude markdown into Word: headings, tables, bold, paragraphs.
    Drops a leading # heading that duplicates the section title already
    added by build_proposal_docx.
    """
    lines = content.replace("\r\n", "\n").split("\n")
    table_buf: list[str] = []
    para_buf: list[str] = []
    skipped_leading_h1 = False

    def flush_table() -> None:
        nonlocal table_buf
        if table_buf:
            _write_table(doc, table_buf)
            table_buf = []

    def flush_para() -> None:
        nonlocal para_buf
        text = " ".join(p.strip() for p in para_buf if p.strip()).strip()
        para_buf = []
        if text:
            p = doc.add_paragraph()
            _add_text_with_bold(p, text)

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        indent = len(line) - len(line.lstrip(" "))
        if (
            not skipped_leading_h1
            and stripped.startswith("# ")
            and not para_buf
            and not table_buf
        ):
            skipped_leading_h1 = True
            continue
        if _is_table_row(stripped):
            flush_para()
            table_buf.append(stripped)
            continue
        if table_buf:
            flush_table()
        if not stripped:
            flush_para()
            continue
        # Horizontal rules are section dividers in markdown; Word already
        # separates sections with headings, so emit nothing rather than a
        # literal "---" paragraph.
        if _HR_RE.match(stripped):
            flush_para()
            continue
        for marker, level in (("#### ", 4), ("### ", 3), ("## ", 2), ("# ", 2)):
            if stripped.startswith(marker):
                flush_para()
                doc.add_heading(stripped[len(marker):].strip(), level=level)
                break
        else:
            bullet = _BULLET_RE.match(stripped)
            if bullet:
                flush_para()
                _add_list_item(doc, bullet.group(1), indent, numbered=False)
                continue
            numbered = _NUMBER_RE.match(stripped)
            if numbered:
                flush_para()
                _add_list_item(doc, numbered.group(1), indent, numbered=True)
                continue
            para_buf.append(stripped)

    flush_table()
    flush_para()


def _insert_toc_field(doc: Document) -> None:
    """
    Inserts a real, clickable, page-numbered Word TOC field. python-docx
    has no built-in TOC helper — this is the standard raw-OXML field-code
    approach. Word computes the actual page numbers when the file is
    opened; some Word versions show "Right-click → Update Field" on
    first open instead of populating instantly — that's normal Word
    behavior, not a bug in this code.

    If this proves fragile in testing, an acceptable fallback is a
    static list of section names with no page numbers, added as plain
    paragraphs instead — flag it rather than spending excess time
    fighting Word's field-code quirks.
    """
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-1" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, fld_end):
        run._r.append(el)


def build_proposal_docx(
    sections: dict,
    opportunity: dict,
    output_dir: str = "output/proposals",
) -> str:
    """
    Builds a formatted .docx from a proposal sections dict. Returns the
    file path. Sections missing from `sections` are skipped silently —
    this must not crash on a partial (lightweight/WATCH-tier) dict.
    """
    os.makedirs(output_dir, exist_ok=True)
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Title page
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("CORTECH CONSULTING GROUP")
    header_run.bold = True
    header_run.font.size = Pt(16)
    header_run.font.color.rgb = RGBColor(0x14, 0x21, 0x3D)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_type = (
        "Expression of Interest"
        if sections.get("submission_type") == "EOI"
        else "Technical Proposal"
    )
    subtitle.add_run(doc_type).italic = True

    doc.add_paragraph()

    opp_title = doc.add_paragraph()
    opp_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = opp_title.add_run(opportunity.get("title", "Untitled Assignment"))
    title_run.bold = True
    title_run.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Client: {opportunity.get('client', 'N/A')}\n")
    meta.add_run("Submitted by: Cortech Consulting Group\n")
    meta.add_run(f"Date: {datetime.now().strftime('%d %B %Y')}")

    doc.add_page_break()
    _insert_toc_field(doc)
    doc.add_page_break()

    sections_included = 0
    for key, heading in SECTION_ORDER:
        content = sections.get(key)
        if not content or not isinstance(content, str):
            continue
        content = _client_safe_section(key, content)
        if not content.strip():
            continue
        doc.add_heading(heading, level=1)
        _write_markdown_content(doc, content)
        doc.add_page_break()
        sections_included += 1

    if sections_included == 0:
        logger.warning("build_proposal_docx: no sections had content — file will be nearly empty")

    safe_title = "".join(
        c for c in opportunity.get("title", "proposal") if c.isalnum() or c in " -_"
    )[:60].strip() or "proposal"
    file_path = os.path.join(output_dir, f"{safe_title}.docx")
    doc.save(file_path)
    logger.success(f"Built proposal docx ({sections_included} sections): {file_path}")
    return file_path
