# processors/downloader.py
import asyncio
import httpx
import pdfplumber
from docx import Document as DocxDocument
from io import BytesIO
from zipfile import ZipFile
from loguru import logger
from database.supabase_client import store_document
import re
import time
from urllib.parse import parse_qs, urlparse, urljoin

from processors.document_quality import MIN_USEFUL_CHARS, assess_extraction
from config import (
    DOCUMENT_DOWNLOAD_TIMEOUT_SECONDS,
    DOCUMENT_DOWNLOAD_MAX_RETRIES,
    MAX_DOCUMENT_BYTES,
    MAX_DOCUMENT_UNCOMPRESSED_BYTES,
    MAX_DOWNLOAD_REDIRECTS,
    MAX_EXTRACTED_TEXT_CHARS,
    MAX_GDRIVE_FILES,
)
from utils.errors import ErrorType
from utils.browser_security import install_browser_request_guard, launch_chromium
from utils.urls import UnsafeURLError, assert_public_http_url, assert_safe_redirect, safe_filename

HTTP_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    )
}


def _looks_like_pdf_url(url: str) -> bool:
    return url.lower().split("?")[0].split("#")[0].endswith(".pdf")


def _retryable_download_error(exc: Exception) -> bool:
    """Retry only failures that can plausibly improve without changing input."""
    if isinstance(exc, (UnsafeURLError, ValueError)):
        return False
    if isinstance(exc, (httpx.TimeoutException, httpx.NetworkError)):
        return True
    if isinstance(exc, httpx.HTTPStatusError):
        status = exc.response.status_code
        return status in {408, 425, 429, 500, 502, 503, 504}
    return False


def _download_document_once(url: str) -> bytes:
    """One secure fetch attempt. Redirect targets are validated per hop."""
    current_url = url
    timeout = httpx.Timeout(DOCUMENT_DOWNLOAD_TIMEOUT_SECONDS)
    with httpx.Client(timeout=timeout, follow_redirects=False) as client:
        for redirect_count in range(MAX_DOWNLOAD_REDIRECTS + 1):
            # Protect every hop, not only the original listing URL.
            assert_public_http_url(current_url, resolve=True)
            with client.stream("GET", current_url, headers=HTTP_HEADERS) as response:
                if response.status_code in {301, 302, 303, 307, 308}:
                    location = response.headers.get("location")
                    if not location:
                        response.raise_for_status()
                        raise RuntimeError("Redirect response did not include a Location header")
                    if redirect_count >= MAX_DOWNLOAD_REDIRECTS:
                        raise RuntimeError(
                            f"Too many redirects (max {MAX_DOWNLOAD_REDIRECTS}) while downloading document"
                        )
                    next_url = urljoin(current_url, location)
                    # Validate now as well as at the next loop's start so an
                    # internal redirect is never accidentally requested.
                    assert_safe_redirect(current_url, next_url)
                    current_url = next_url
                    continue

                response.raise_for_status()
                declared_size = response.headers.get("content-length")
                if declared_size:
                    try:
                        if int(declared_size) > MAX_DOCUMENT_BYTES:
                            raise ValueError(
                                f"Document exceeds {MAX_DOCUMENT_BYTES} byte limit "
                                f"(declared {declared_size} bytes)"
                            )
                    except ValueError as exc:
                        # A malformed Content-Length is not a security signal;
                        # stream accounting below remains authoritative. Keep
                        # a real over-limit error visible to the caller.
                        if "exceeds" in str(exc):
                            raise

                content = bytearray()
                for chunk in response.iter_bytes():
                    content.extend(chunk)
                    if len(content) > MAX_DOCUMENT_BYTES:
                        raise ValueError(
                            f"Document exceeds {MAX_DOCUMENT_BYTES} byte limit while streaming"
                        )
                return bytes(content)

    # The loop either returns, raises, or reaches this impossible guard.
    raise RuntimeError("Document download ended without a response")


def download_document(url: str) -> bytes:
    """Download an untrusted document within explicit safety limits.

    httpx does not re-run our URL policy for redirects. Redirects therefore
    have to be followed manually, with every next hop checked before a
    request is made. TLS verification intentionally remains enabled: a bad
    certificate is a failed source, not a reason to weaken transport
    security for procurement documents.
    """
    try:
        assert_public_http_url(url, resolve=True)
    except UnsafeURLError as e:
        logger.error(f"Download blocked ({ErrorType.SSRF_ERROR}): {e}")
        raise

    for attempt in range(DOCUMENT_DOWNLOAD_MAX_RETRIES + 1):
        try:
            return _download_document_once(url)
        except Exception as exc:
            if not _retryable_download_error(exc) or attempt >= DOCUMENT_DOWNLOAD_MAX_RETRIES:
                logger.error(
                    f"Download failed for {url}: {exc} "
                    f"error_type={ErrorType.INGESTION_ERROR}"
                )
                raise
            # 1s, 2s, then give up. This is intentionally capped so a bad
            # source cannot hold up the entire discovery run for minutes.
            delay = 2 ** attempt
            logger.warning(
                f"Transient download failure (attempt {attempt + 1}/"
                f"{DOCUMENT_DOWNLOAD_MAX_RETRIES + 1}); retrying in {delay}s: {exc}"
            )
            time.sleep(delay)

    raise RuntimeError("Document download retry loop ended unexpectedly")


def extract_text_from_pdf(content: bytes) -> str:
    """Extract text from PDF bytes."""
    try:
        with pdfplumber.open(BytesIO(content)) as pdf:
            pages = []
            chars = 0
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    remaining = MAX_EXTRACTED_TEXT_CHARS - chars
                    if remaining <= 0:
                        break
                    pages.append(text[:remaining])
                    chars += len(pages[-1])
                    if chars >= MAX_EXTRACTED_TEXT_CHARS:
                        logger.warning("PDF extracted text reached configured character cap")
                        break
            return "\n\n".join(pages)[:MAX_EXTRACTED_TEXT_CHARS]
    except Exception as e:
        logger.error(f"PDF extraction failed: {e} error_type={ErrorType.DOCUMENT_ERROR}")
        return ""


def extract_text_from_docx(content: bytes) -> str:
    """Extract text from DOCX bytes."""
    try:
        # DOCX is a ZIP container. A 25 MiB download can otherwise inflate to
        # gigabytes of XML before python-docx reaches our text-character cap.
        with ZipFile(BytesIO(content)) as archive:
            uncompressed = sum(info.file_size for info in archive.infolist())
        if uncompressed > MAX_DOCUMENT_UNCOMPRESSED_BYTES:
            raise ValueError(
                "DOCX uncompressed content exceeds configured safety limit "
                f"({uncompressed} bytes)"
            )
        doc = DocxDocument(BytesIO(content))
        parts = []
        chars = 0

        def append_text(value: str) -> bool:
            nonlocal chars
            if not value or not value.strip() or chars >= MAX_EXTRACTED_TEXT_CHARS:
                return chars < MAX_EXTRACTED_TEXT_CHARS
            remaining = MAX_EXTRACTED_TEXT_CHARS - chars
            parts.append(value[:remaining])
            chars += len(parts[-1])
            return chars < MAX_EXTRACTED_TEXT_CHARS

        for paragraph in doc.paragraphs:
            if not append_text(paragraph.text):
                break
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if not append_text(row_text):
                    break
            if chars >= MAX_EXTRACTED_TEXT_CHARS:
                break

        if chars >= MAX_EXTRACTED_TEXT_CHARS:
            logger.warning("DOCX extracted text reached configured character cap")
        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e} error_type={ErrorType.DOCUMENT_ERROR}")
        return ""


def _find_pdf_links(html: str, base_url: str) -> list[str]:
    """Collect PDF hrefs from a rendered page — common on WordPress tender posts."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    seen: set[str] = set()
    links: list[str] = []
    for anchor in soup.find_all("a", href=True):
        href = anchor["href"]
        if ".pdf" not in href.lower():
            continue
        absolute = urljoin(base_url, href)
        if absolute not in seen:
            seen.add(absolute)
            links.append(absolute)
    return links


def _extract_from_pdf_url(pdf_url: str) -> tuple[str, bytes]:
    """Download and extract text from a direct PDF link."""
    content = download_document(pdf_url)
    return extract_text_from_pdf(content), content


def extract_text_from_html(html: str) -> str:
    """Extract clean text from HTML."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)[:MAX_EXTRACTED_TEXT_CHARS]


async def _fetch_rendered_html(url: str) -> str | None:
    """
    Playwright fallback for individual ToR/opportunity pages that render
    via JS. Same problem monitors/scraper.py already solves for listing
    pages, one layer deeper — the linked document page itself.
    """
    from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout

    try:
        assert_public_http_url(url, resolve=True)
    except UnsafeURLError as e:
        logger.error(f"  Playwright blocked ({ErrorType.SSRF_ERROR}): {e}")
        return None

    try:
        async with async_playwright() as p:
            browser = await launch_chromium(p)
            # Do not weaken TLS for untrusted tender sources. A bad certificate
            # is a failed source, not a reason to accept tampered content.
            context = await browser.new_context()
            page = await context.new_page()
            try:
                await install_browser_request_guard(page)
                await page.goto(url, wait_until="commit", timeout=45000)
                await page.wait_for_timeout(5000)
                return await page.content()
            except PlaywrightTimeout:
                logger.warning(f"  Playwright timeout loading: {url}")
                return None
            finally:
                await page.close()
                await context.close()
                await browser.close()
    except Exception as e:
        logger.error(f"  Playwright fallback failed for {url}: {e}")
        return None


def _fetch_rendered_html_sync(url: str) -> str | None:
    """Sync wrapper — same pattern as scrape_non_rss_sources() in scraper.py."""
    try:
        return asyncio.run(_fetch_rendered_html(url))
    except Exception as e:
        logger.error(f"  Browser fallback crashed for {url}: {e}")
        return None


# ── GOOGLE DRIVE (folder + multi-file annex packs) ───────────────────────────

_GDRIVE_HOSTS = ("drive.google.com", "docs.google.com")


def _is_gdrive_url(url: str) -> bool:
    host = urlparse(url).netloc.lower()
    return any(h in host for h in _GDRIVE_HOSTS)


def _gdrive_id_from_url(url: str) -> tuple[str, str]:
    """
    Returns (kind, id) where kind is 'folder' or 'file'.
    kind='file' covers uploaded PDFs/DOCX and native Google Docs.
    """
    parsed = urlparse(url)
    path = parsed.path or ""
    query = parse_qs(parsed.query)

    folder_match = re.search(r"/folders/([a-zA-Z0-9_-]+)", path)
    if folder_match:
        return "folder", folder_match.group(1)

    file_match = re.search(r"/file/d/([a-zA-Z0-9_-]+)", path)
    if file_match:
        return "file", file_match.group(1)

    doc_match = re.search(r"/document/d/([a-zA-Z0-9_-]+)", path)
    if doc_match:
        return "file", doc_match.group(1)

    if query.get("id"):
        file_id = query["id"][0]
        if "/folders" in path:
            return "folder", file_id
        return "file", file_id

    return "file", ""


def _extract_from_bytes(content: bytes, filename: str = "") -> tuple[str, str]:
    """Detect PDF/DOCX/HTML from bytes or filename and extract text."""
    name = filename.lower()
    if content[:4] == b"%PDF" or name.endswith(".pdf"):
        return extract_text_from_pdf(content), "pdf"
    if content[:2] == b"PK" or name.endswith(".docx"):
        return extract_text_from_docx(content), "docx"
    try:
        return extract_text_from_html(content.decode("utf-8", errors="ignore")), "html"
    except Exception:
        return "", "bin"


def _gdrive_confirm_url(html: str, file_id: str) -> str | None:
    """Virus-scan interstitial — follow the confirm= token if present."""
    match = re.search(
        rf"href=\"(/uc\?export=download[^\"']*confirm=[^\"']+id={re.escape(file_id)}[^\"']*)\"",
        html,
        re.I,
    )
    if match:
        return "https://drive.google.com" + match.group(1).replace("&amp;", "&")
    match = re.search(r"confirm=([0-9A-Za-z_-]+)", html)
    if match:
        return (
            f"https://drive.google.com/uc?export=download"
            f"&confirm={match.group(1)}&id={file_id}"
        )
    return None


def _download_gdrive_file(file_id: str) -> bytes:
    """Download a publicly-shared Drive file (Anyone with the link)."""
    download_url = f"https://drive.google.com/uc?export=download&id={file_id}&confirm=t"
    content = download_document(download_url)

    looks_binary = content[:4] == b"%PDF" or content[:2] == b"PK"
    if looks_binary:
        return content

    html = content.decode("utf-8", errors="ignore")
    if "virus scan" in html.lower() or "confirm=" in html:
        confirm_url = _gdrive_confirm_url(html, file_id)
        if confirm_url:
            content = download_document(confirm_url)
            if content[:4] == b"%PDF" or content[:2] == b"PK":
                return content

    # Native Google Doc — export as DOCX, then PDF. Never return the HTML
    # wrapper: a 401/sign-in page was being extracted as a 200-char "ToR".
    for fmt, magic in (("docx", b"PK"), ("pdf", b"%PDF")):
        export_url = (
            f"https://docs.google.com/document/d/{file_id}/export?format={fmt}"
        )
        try:
            exported = download_document(export_url)
            if exported[: len(magic)] == magic:
                return exported
        except Exception as e:
            logger.warning(
                f"  Google Doc {fmt} export failed for {file_id[:12]}: {e}"
            )

    raise RuntimeError(
        f"Could not download Drive file {file_id[:12]} as a document "
        "(got a login/access page, not a PDF or DOCX). Share it as "
        "'Anyone with the link can view' and retry."
    )


def _list_gdrive_folder_files(folder_id: str) -> list[tuple[str, str]]:
    """
    List files in a public Drive folder via the embed view.
    Returns [(file_id, filename), ...] sorted so Annex I/II/III stay in order.
    """
    embed_url = f"https://drive.google.com/embeddedfolderview?id={folder_id}"
    html = ""
    try:
        html = download_document(embed_url).decode("utf-8", errors="ignore")
    except Exception as e:
        logger.warning(f"  Drive folder embed fetch failed: {e}")

    if len(html) < 200:
        rendered = _fetch_rendered_html_sync(
            f"https://drive.google.com/drive/folders/{folder_id}"
        )
        html = rendered or html

    files: list[tuple[str, str]] = []
    seen: set[str] = set()

    for match in re.finditer(
        r"https://drive\.google\.com/file/d/([a-zA-Z0-9_-]+)/[^\"'\s]*\"[^>]*>([^<]+)",
        html,
    ):
        file_id, name = match.group(1), match.group(2).strip()
        if file_id not in seen and name:
            seen.add(file_id)
            files.append((file_id, name))

    if not files:
        for match in re.finditer(
            r"/file/d/([a-zA-Z0-9_-]+)",
            html,
        ):
            file_id = match.group(1)
            if file_id not in seen:
                seen.add(file_id)
                files.append((file_id, f"file_{file_id[:8]}"))

    files.sort(key=lambda item: item[1].lower())
    if len(files) > MAX_GDRIVE_FILES:
        logger.warning(
            f"Drive folder has {len(files)} files; limiting this opportunity to "
            f"the first {MAX_GDRIVE_FILES}"
        )
    return files[:MAX_GDRIVE_FILES]


def fetch_gdrive_and_extract(url: str) -> str:
    """
    Download every file in a public Google Drive folder (or a single file)
    and concatenate extracted text. Each file is labelled so Claude can
    see Annex I / II / III as separate source documents.
    """
    kind, drive_id = _gdrive_id_from_url(url)
    if not drive_id:
        logger.error(f"  Could not parse Google Drive id from: {url}")
        return ""

    items: list[tuple[str, str]] = []
    if kind == "folder":
        logger.info(f"  Google Drive folder — listing files...")
        items = _list_gdrive_folder_files(drive_id)
        logger.info(f"  Found {len(items)} file(s) in Drive folder")
    else:
        items = [(drive_id, "drive_file")]

    if not items:
        logger.error(
            "  Drive folder listed 0 files. Share it as "
            "'Anyone with the link can view' and retry."
        )
        return ""

    parts: list[str] = []
    total_chars = 0
    for file_id, filename in items:
        logger.info(f"  Downloading Drive file: {filename}")
        try:
            content = _download_gdrive_file(file_id)
        except Exception as e:
            logger.warning(f"  Failed to download '{filename}': {e}")
            continue
        text, file_type = _extract_from_bytes(content, filename)
        if not text or len(text) < 40:
            logger.warning(
                f"  No useful text from '{filename}' ({file_type}, {len(content)} bytes)"
            )
            continue
        logger.success(f"  Extracted {len(text):,} chars from {filename}")
        remaining = MAX_EXTRACTED_TEXT_CHARS - total_chars
        if remaining <= 0:
            logger.warning("Drive annex pack reached configured combined text cap")
            break
        part = f"\n\n===== SOURCE FILE: {filename} =====\n\n{text.strip()}\n"
        parts.append(part[:remaining])
        total_chars += len(parts[-1])

    combined = "".join(parts).strip()
    if combined:
        logger.success(
            f"  Combined {len(parts)} Drive file(s) into {len(combined):,} chars"
        )
    return combined


def fetch_and_extract(url: str, opportunity_id: str = None) -> str:
    """
    Main function: fetch URL, detect type, extract text.
    Falls back to a Playwright-rendered page when the plain httpx result
    is too thin to be a real ToR — the storage side-effect is isolated
    from the extraction result so a Supabase hiccup can never discard a
    successful extraction.
    """
    try:
        assert_public_http_url(url, resolve=True)
    except UnsafeURLError as e:
        logger.error(f"Fetch blocked ({ErrorType.SSRF_ERROR}): {e}")
        return ""

    if _is_gdrive_url(url):
        text = fetch_gdrive_and_extract(url)
        quality = assess_extraction(text, source=url)
        if text and not quality["ok"]:
            logger.error(
                f"  {quality['reason']} error_type={quality['error_type']}"
            )
            return ""
        if text:
            logger.success(f"Extracted {len(text)} chars from gdrive: {url[:60]}...")
        else:
            logger.warning(
                f"No text extracted for {url[:60]}... error_type={ErrorType.DOCUMENT_ERROR}"
            )
        return text

    text = ""
    file_type = "html"
    content = b""
    is_pdf_url = _looks_like_pdf_url(url)

    try:
        content = download_document(url)
        if is_pdf_url or content[:4] == b"%PDF":
            text = extract_text_from_pdf(content)
            file_type = "pdf"
        elif url.lower().split("?")[0].endswith(".docx"):
            text = extract_text_from_docx(content)
            file_type = "docx"
        else:
            text = extract_text_from_html(content.decode("utf-8", errors="ignore"))
            file_type = "html"
    except Exception as e:
        if is_pdf_url:
            logger.error(f"  PDF download failed for {url}: {e}")
        else:
            logger.warning(f"  httpx fetch/extract failed for {url}: {e} — will try Playwright")

    # ── PLAYWRIGHT FALLBACK (HTML pages only — never for direct PDF links) ─
    rendered_html: str | None = None
    if file_type == "html" and not is_pdf_url and len(text) < MIN_USEFUL_CHARS:
        logger.info(
            f"  httpx extraction thin ({len(text)} chars) — trying Playwright: {url[:60]}"
        )
        rendered_html = _fetch_rendered_html_sync(url)
        if rendered_html:
            browser_text = extract_text_from_html(rendered_html)
            if len(browser_text) > len(text):
                text = browser_text
                logger.success(
                    f"  Playwright fallback recovered {len(text)} chars: {url[:60]}"
                )

    # ── PDF LINK FOLLOW — WordPress posts often only link to the ToR PDF ───
    html_source = rendered_html or (
        content.decode("utf-8", errors="ignore") if content else ""
    )
    if file_type == "html" and len(text) < MIN_USEFUL_CHARS and html_source:
        for pdf_url in _find_pdf_links(html_source, url):
            try:
                assert_public_http_url(pdf_url, resolve=True)
            except UnsafeURLError as e:
                logger.warning(f"  Skipping PDF link ({ErrorType.SSRF_ERROR}): {e}")
                continue
            logger.info(f"  Following PDF link from page: {pdf_url[:70]}")
            try:
                pdf_text, pdf_content = _extract_from_pdf_url(pdf_url)
            except Exception as e:
                logger.warning(f"  PDF link fetch failed ({pdf_url[:60]}): {e}")
                continue
            if len(pdf_text) >= MIN_USEFUL_CHARS:
                text = pdf_text
                content = pdf_content
                file_type = "pdf"
                url = pdf_url  # store under the PDF filename
                logger.success(
                    f"  Recovered {len(text)} chars from linked PDF: {pdf_url[:60]}"
                )
                break

    # ── STORE IN SUPABASE — isolated, never discards a good extraction ──
    if opportunity_id and file_type in ["pdf", "docx"] and text and content:
        try:
            file_name = safe_filename(
                url.split("/")[-1] or f"document.{file_type}",
                default=f"document.{file_type}",
            )
            store_document(opportunity_id, file_name, content, file_type)
        except Exception as e:
            logger.warning(
                f"  Supabase document store failed (non-fatal): {e} "
                f"error_type={ErrorType.STORAGE_ERROR}"
            )

    quality = assess_extraction(text, source=url)
    if text and not quality["ok"]:
        logger.error(
            f"  {quality['reason']} error_type={quality['error_type']}"
        )
        return ""

    if text:
        logger.success(f"Extracted {len(text)} chars from {file_type}: {url[:60]}...")
    else:
        logger.warning(
            f"No text extracted for {url[:60]}... error_type={ErrorType.DOCUMENT_ERROR}"
        )
    return text
