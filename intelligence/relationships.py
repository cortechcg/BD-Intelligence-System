"""Cited relationship edges from Cortech past submissions (Phase 5).

Deterministic patterns only. A partner name not in the excerpt is not stored.
See docs/adr/009-competitor-relationship-intelligence.md.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from intelligence.organizations import (
    MAX_ORG_NAME_CHARS,
    MatchResult,
    match_organization,
    sanitize_org_name,
)
from utils.untrusted import wrap_untrusted

EVIDENCE_VERIFIED = "VERIFIED"
MAX_TEXT_CHARS = 400_000
MAX_EXCERPT_CHARS = 500
MAX_WINDOW = 900

KIND_JV = "joint_venture"
KIND_CONSORTIUM = "consortium"
KIND_SUB = "subcontractor"
KIND_ASSOC = "in_association"

_CORTECH = re.compile(r"\bcortech\b", re.I)
_SELF_NAME = re.compile(
    r"^cortech(\s+consulting(\s+(group|llc|ltd|limited))?)?(\s+(llc|ltd|limited|group))?$",
    re.I,
)
_NO_PARTNER = re.compile(
    r"no jv partner|sole firm with no jv|with no jv partner and no sub-consultants",
    re.I,
)
_TYPICAL = re.compile(
    r"typically work with|leading un agencies|leading development agencies|"
    r"leading international|we have also partnered with leading",
    re.I,
)
_POLICY_NOISE = re.compile(
    r"sub-?contractors to adopt|sub-contractor or consortium member|"
    r"contracts and partnership agreements include a standard clause|"
    r"if the partner notes that it does not have subcontractors",
    re.I,
)
_PERSON_ROLE = re.compile(
    r"^(?:dr|mr|mrs|ms|prof)\b|lead consultant|co-consultant|"
    r"market research officer|agronomist|water resource specialist",
    re.I,
)
_COUNTRY_CODE = re.compile(r"\s+\([A-Z]{2}\)\s*$")
_TRAILING_JV = re.compile(
    r"\s+(?:the consortium or the consultant|together hereinafter.*)$",
    re.I,
)

_JV_WITH = re.compile(
    r"in joint venture with\s+(.+?)(?:"
    r"\s+\([A-Z]{2}\)\s*(?:\(together hereinafter|\s+for the provision|$)|"
    r"\s+\(together hereinafter|"
    r"\s+for the provision|"
    r"\s*\.(?:\s|$)"
    r")",
    re.I | re.S,
)
_LEADER = re.compile(
    r"([A-Z][A-Za-z0-9 .,&'/–\-]{1,80}?)\s+as Leader of the Consortium "
    r"and\s+Cortech",
    re.I,
)
_COMMISSIONED = re.compile(
    r"in partnership with\s+([A-Z][A-Za-z0-9 .,&'/–\-]{1,80}?)"
    r"\s+and commissioned by",
    re.I,
)
_IN_ASSOC = re.compile(
    r"in association with\s+([A-Z][A-Za-z0-9 .,&'/–\-]{1,80}?)"
    r"(?:\s+for\s|\s+to\s|\s*\.|$)",
    re.I,
)
_SUB_WITH = re.compile(
    r"sub-?contract(?:or|ed|ing)?\s+(?:agreement\s+)?"
    r"(?:with|to)\s+([A-Z][A-Za-z0-9 .,&'/–\-]{1,80}?)"
    r"(?:\s+for\s|\s*\.|$)",
    re.I,
)
_NAMED_FIRM_INTRO = re.compile(
    r"([A-Z][A-Za-z][A-Za-z0-9 .,&'/–\-]{2,80}?)\s+\(([A-Z]{2,8})\)\s+is a\b",
)
_MARKER = re.compile(
    r"joint venture|consortium|in association with|sub-?contract|commissioned by",
    re.I,
)


@dataclass(frozen=True)
class RelationshipEdge:
    observed_name: str
    relationship_kind: str
    excerpt: str
    document_name: str
    chunk_id: str
    cortech_role: str = ""
    evidence_status: str = EVIDENCE_VERIFIED
    match: MatchResult | None = None

    def as_dict(self) -> dict:
        payload = {
            "observed_name": self.observed_name,
            "relationship_kind": self.relationship_kind,
            "excerpt": self.excerpt,
            "document_name": self.document_name,
            "chunk_id": self.chunk_id,
            "cortech_role": self.cortech_role,
            "evidence_status": self.evidence_status,
            "citation": f"{self.document_name} ({self.chunk_id})",
        }
        if self.match is not None:
            payload["match"] = self.match.as_dict()
        return payload


def wrap_proposal_text(text: str) -> str:
    return wrap_untrusted((text or "")[:MAX_TEXT_CHARS])


def _norm_contains(needle: str, haystack: str) -> bool:
    n = re.sub(r"[^a-z0-9]+", " ", (needle or "").casefold()).strip()
    h = re.sub(r"[^a-z0-9]+", " ", (haystack or "").casefold()).strip()
    return bool(n) and n in h


def _clean_partner(raw: str) -> str:
    text = sanitize_org_name(raw or "")
    text = _COUNTRY_CODE.sub("", text).strip(" ,;|-")
    text = _TRAILING_JV.sub("", text).strip(" ,;|-")
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) > MAX_ORG_NAME_CHARS:
        text = text[:MAX_ORG_NAME_CHARS].rstrip()
    if not text or _SELF_NAME.match(text):
        return ""
    if _PERSON_ROLE.search(text):
        return ""
    if len(text) < 4:
        return ""
    # Require at least one letter; reject table junk.
    if not re.search(r"[A-Za-z]", text):
        return ""
    return text


def cited_edge_or_none(
    name: str,
    excerpt: str,
    *,
    document_name: str,
    chunk_id: str,
    relationship_kind: str,
    cortech_role: str = "",
) -> RelationshipEdge | None:
    cleaned = _clean_partner(name)
    snippet = " ".join((excerpt or "").split())
    if len(snippet) > MAX_EXCERPT_CHARS:
        snippet = snippet[:MAX_EXCERPT_CHARS].rstrip()
    doc = sanitize_org_name(document_name) or (document_name or "").strip()
    if not cleaned or not snippet or not doc:
        return None
    if not _norm_contains(cleaned, snippet):
        return None
    if not _CORTECH.search(snippet):
        return None
    return RelationshipEdge(
        observed_name=cleaned,
        relationship_kind=relationship_kind,
        excerpt=snippet,
        document_name=doc[:240],
        chunk_id=(chunk_id or "")[:240] or f"proposal-file:{doc}:excerpt",
        cortech_role=cortech_role,
        evidence_status=EVIDENCE_VERIFIED,
    )


def _window(text: str, start: int, end: int) -> str:
    lo = max(0, start - 80)
    hi = min(len(text), end + MAX_WINDOW)
    return text[lo:hi]


def _add_edge(
    out: list[RelationshipEdge],
    seen: set[tuple[str, str]],
    edge: RelationshipEdge | None,
) -> None:
    if edge is None:
        return
    key = (edge.observed_name.casefold(), edge.relationship_kind)
    if key in seen:
        return
    seen.add(key)
    out.append(edge)


def extract_relationship_edges(
    text: str,
    *,
    document_name: str,
) -> list[RelationshipEdge]:
    """Extract cited Cortech–counterpart edges. Empty if the file names nobody."""
    wrap_proposal_text(text)
    raw = (text or "")[:MAX_TEXT_CHARS]
    if not raw.strip() or not document_name:
        return []

    out: list[RelationshipEdge] = []
    seen: set[tuple[str, str]] = set()
    doc = Path(document_name).name

    for match in _MARKER.finditer(raw):
        window = _window(raw, match.start(), match.end())
        if not _CORTECH.search(window):
            continue
        if _NO_PARTNER.search(window):
            continue
        if _TYPICAL.search(window):
            continue
        if _POLICY_NOISE.search(window):
            continue
        offset = match.start()
        chunk_id = f"proposal-file:{doc}:{offset}"

        jv = _JV_WITH.search(window)
        if jv:
            _add_edge(
                out,
                seen,
                cited_edge_or_none(
                    jv.group(1),
                    window,
                    document_name=doc,
                    chunk_id=chunk_id,
                    relationship_kind=KIND_JV,
                    cortech_role="member",
                ),
            )

        leader = _LEADER.search(window)
        if leader:
            _add_edge(
                out,
                seen,
                cited_edge_or_none(
                    leader.group(1),
                    window,
                    document_name=doc,
                    chunk_id=chunk_id,
                    relationship_kind=KIND_CONSORTIUM,
                    cortech_role="partner",
                ),
            )

        commissioned = _COMMISSIONED.search(window)
        if commissioned:
            _add_edge(
                out,
                seen,
                cited_edge_or_none(
                    commissioned.group(1),
                    window,
                    document_name=doc,
                    chunk_id=chunk_id,
                    relationship_kind=KIND_ASSOC,
                    cortech_role="joint_delivery",
                ),
            )

        assoc = _IN_ASSOC.search(window)
        if assoc:
            _add_edge(
                out,
                seen,
                cited_edge_or_none(
                    assoc.group(1),
                    window,
                    document_name=doc,
                    chunk_id=chunk_id,
                    relationship_kind=KIND_ASSOC,
                    cortech_role="member",
                ),
            )

        sub = _SUB_WITH.search(window)
        if sub:
            _add_edge(
                out,
                seen,
                cited_edge_or_none(
                    sub.group(1),
                    window,
                    document_name=doc,
                    chunk_id=chunk_id,
                    relationship_kind=KIND_SUB,
                    cortech_role="lead",
                ),
            )

        if re.search(r"joint venture", window, re.I):
            intro = _NAMED_FIRM_INTRO.search(window)
            if intro:
                _add_edge(
                    out,
                    seen,
                    cited_edge_or_none(
                        intro.group(1),
                        window,
                        document_name=doc,
                        chunk_id=chunk_id,
                        relationship_kind=KIND_JV,
                        cortech_role="lead",
                    ),
                )

    return out


def bind_relationship_edges(
    edges: Sequence[RelationshipEdge],
    index: Sequence | None = None,
) -> list[RelationshipEdge]:
    from intelligence.organizations import index_with_match

    working = list(index or ())
    bound: list[RelationshipEdge] = []
    for edge in edges:
        match = match_organization(edge.observed_name, working)
        working = index_with_match(working, match, entity_kind="unknown")
        bound.append(
            RelationshipEdge(
                observed_name=edge.observed_name,
                relationship_kind=edge.relationship_kind,
                excerpt=edge.excerpt,
                document_name=edge.document_name,
                chunk_id=edge.chunk_id,
                cortech_role=edge.cortech_role,
                evidence_status=edge.evidence_status,
                match=match,
            )
        )
    return bound


def read_proposal_file(path: Path) -> str:
    """Local PDF/DOCX/text reader. Missing/unreadable files return empty."""
    try:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            import pdfplumber

            parts: list[str] = []
            with pdfplumber.open(str(path)) as pdf:
                for page in pdf.pages[:12]:
                    text = page.extract_text() or ""
                    if text.strip():
                        parts.append(text)
            return "\n".join(parts)[:MAX_TEXT_CHARS]
        if suffix == ".docx":
            from docx import Document as DocxDocument

            doc = DocxDocument(str(path))
            sections = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        sections.append(" | ".join(cells))
            return "\n".join(sections)[:MAX_TEXT_CHARS]
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="ignore")[:MAX_TEXT_CHARS]
    except Exception:
        return ""
    return ""


def extract_from_proposal_path(path: Path | str) -> list[RelationshipEdge]:
    target = Path(path)
    if not target.is_file() or target.name.startswith("~$"):
        return []
    if target.name.startswith("PUT_"):
        return []
    text = read_proposal_file(target)
    return extract_relationship_edges(text, document_name=target.name)


def extract_from_proposals_dir(
    directory: Path | str | None = None,
    *,
    persist: bool = False,
    index: Sequence | None = None,
) -> list[RelationshipEdge]:
    root = Path(directory) if directory else Path("data/proposals")
    if not root.is_dir():
        return []
    from intelligence.organizations import index_with_match

    working = list(index or ())
    all_edges: list[RelationshipEdge] = []
    for path in sorted(root.iterdir()):
        if path.suffix.lower() not in {".pdf", ".docx", ".txt", ".md"}:
            continue
        edges = extract_from_proposal_path(path)
        bound = bind_relationship_edges(edges, working)
        for edge in bound:
            if edge.match is not None:
                working = index_with_match(
                    working, edge.match, entity_kind="unknown"
                )
            if persist:
                try:
                    from database import intelligence_facts as facts_store

                    facts_store.persist_relationship_edge(edge)
                except Exception:
                    pass
            all_edges.append(edge)
    return all_edges


def extract_from_proposal_embeddings(
    rows: Sequence[dict] | None,
    *,
    persist: bool = False,
    index: Sequence | None = None,
) -> list[RelationshipEdge]:
    """Scan already-stored proposal chunks. Missing/malformed rows skip."""
    from intelligence.organizations import index_with_match

    working = list(index or ())
    all_edges: list[RelationshipEdge] = []
    for row in rows or ():
        if not isinstance(row, dict):
            continue
        chunk = row.get("content_chunk")
        if not isinstance(chunk, str) or not chunk.strip():
            continue
        pid = str(row.get("airtable_proposal_id") or row.get("id") or "")
        title = str(row.get("project_title") or pid or "proposal_embeddings")
        meta = row.get("metadata") if isinstance(row.get("metadata"), dict) else {}
        source_file = str(meta.get("source_file") or title)
        edges = extract_relationship_edges(chunk, document_name=source_file)
        remapped = []
        for edge in edges:
            remapped.append(
                RelationshipEdge(
                    observed_name=edge.observed_name,
                    relationship_kind=edge.relationship_kind,
                    excerpt=edge.excerpt,
                    document_name=edge.document_name,
                    chunk_id=f"proposal_embeddings:{pid or edge.chunk_id}",
                    cortech_role=edge.cortech_role,
                    evidence_status=edge.evidence_status,
                )
            )
        bound = bind_relationship_edges(remapped, working)
        for edge in bound:
            if edge.match is not None:
                working = index_with_match(
                    working, edge.match, entity_kind="unknown"
                )
            if persist:
                try:
                    from database import intelligence_facts as facts_store

                    facts_store.persist_relationship_edge(edge)
                except Exception:
                    pass
            all_edges.append(edge)
    return all_edges
