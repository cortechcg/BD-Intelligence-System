"""
╔══════════════════════════════════════════════════════════════════╗
║              CORTECH CV EMBEDDER SCRIPT                          ║
║                                                                  ║
║  Run this AFTER populate_airtable.py                             ║
║  Reads consultants and proposal files and loads them into the    ║
║  Supabase vector store for semantic search matching              ║
║                                                                  ║
║  Run: python embed_cvs.py                                        ║
╚══════════════════════════════════════════════════════════════════╝
"""

import argparse
import errno
import os
import socket
import ssl
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Callable

import httpx
from dotenv import load_dotenv
from rich.console import Console
from rich.progress import track
from loguru import logger

load_dotenv()
console = Console()

EMBEDDING_MODEL = "text-embedding-3-small"
MIN_CV_TEXT_CHARS = 50
MIN_PROPOSAL_TEXT_CHARS = 100

# Per-record retries after a transient error (same consultant, not the next one).
TRANSIENT_BACKOFFS = (5.0, 15.0, 45.0)
CIRCUIT_BREAKER_THRESHOLD = 3
CIRCUIT_BREAKER_PAUSE_SECONDS = 45.0
# urllib3 must not auto-retry 429.
AIRTABLE_RETRY_STATUS_FORCELIST = (500, 502, 503, 504)

_TRANSIENT_MESSAGE_MARKERS = (
    "temporary failure in name resolution",
    "name or service not known",
    "nodename nor servname provided",
    "getaddrinfo failed",
    "handshake operation timed out",
    "the handshake operation timed out",
    "ssl handshake timeout",
    "connection timed out",
    "connect timeout",
    "read timeout",
    "write timeout",
    "pool timeout",
    "too many requests",
    "error code: 429",
    "status code 429",
    "429 too many requests",
    "errno -3",
    "[errno -3]",
)


def _walk_exceptions(exc: BaseException):
    """Yield exc and nested cause/context/reason exceptions without loops."""
    seen: set[int] = set()
    stack: list[BaseException] = [exc]
    while stack:
        current = stack.pop()
        if current is None or id(current) in seen:
            continue
        seen.add(id(current))
        yield current
        if current.__cause__ is not None:
            stack.append(current.__cause__)
        if current.__context__ is not None:
            stack.append(current.__context__)
        reason = getattr(current, "reason", None)
        if isinstance(reason, BaseException):
            stack.append(reason)
        original = getattr(current, "original_error", None)
        if isinstance(original, BaseException):
            stack.append(original)
        for arg in getattr(current, "args", ()):
            if isinstance(arg, BaseException):
                stack.append(arg)


def _http_status_code(exc: BaseException) -> int | None:
    for attr in ("status_code", "status"):
        value = getattr(exc, attr, None)
        if isinstance(value, int) and value > 0:
            return value
    response = getattr(exc, "response", None)
    if response is not None:
        value = getattr(response, "status_code", None)
        if isinstance(value, int):
            return value
    return None


def _is_transient_one(exc: BaseException) -> bool:
    status = _http_status_code(exc)
    if status == 429:
        return True
    if status is not None and 400 <= status < 500:
        return False

    if isinstance(exc, ssl.SSLCertVerificationError):
        return False

    if isinstance(exc, (httpx.TimeoutException, httpx.ConnectError, httpx.NetworkError)):
        return True

    if isinstance(exc, socket.gaierror):
        return True

    if isinstance(exc, socket.timeout):
        return True

    if isinstance(exc, TimeoutError):
        return True

    if isinstance(exc, ConnectionError):
        return True

    if isinstance(exc, ssl.SSLError):
        msg = str(exc).lower()
        return any(token in msg for token in ("timed out", "timeout", "handshake"))

    if isinstance(exc, OSError) and getattr(exc, "errno", None) in (
        -3,
        getattr(errno, "EAI_AGAIN", -3),
        errno.ETIMEDOUT,
        errno.ECONNRESET,
        errno.ENETUNREACH,
        errno.EHOSTUNREACH,
    ):
        return True

    try:
        import requests

        if isinstance(exc, (requests.exceptions.ConnectionError, requests.exceptions.Timeout)):
            return True
    except ImportError:
        pass

    msg = str(exc).lower()
    return any(marker in msg for marker in _TRANSIENT_MESSAGE_MARKERS)


def is_transient_network_error(exc: BaseException) -> bool:
    """True for DNS/SSL/timeout/429 blips; False for auth, schema, and other bugs."""
    top_status = _http_status_code(exc)
    if top_status == 429:
        return True
    if top_status is not None and 400 <= top_status < 500:
        return False
    return any(_is_transient_one(err) for err in _walk_exceptions(exc))


def _format_wait_seconds(seconds: float) -> str:
    if seconds == int(seconds):
        return str(int(seconds))
    return str(seconds)


def announce_network_wait(seconds: float, label: str) -> None:
    msg = f"Network error — waiting {_format_wait_seconds(seconds)}s then retrying {label}"
    console.print(f"[yellow]{msg}[/yellow]")
    logger.warning(msg)


def retry_on_transient(
    operation: Callable,
    *,
    label: str,
    backoffs: tuple[float, ...] = TRANSIENT_BACKOFFS,
    sleep_fn=None,
):
    """
    Run operation(); on a transient network error wait and retry the same work.

    Permanent errors raise immediately. Exhausted retries re-raise the last
    transient error so the caller can count a failure.
    """
    if sleep_fn is None:
        sleep_fn = time.sleep
    last_transient: Exception | None = None
    attempts = (0.0,) + backoffs
    for attempt, wait in enumerate(attempts):
        if wait:
            announce_network_wait(wait, label)
            sleep_fn(wait)
        try:
            return operation()
        except Exception as exc:
            if not is_transient_network_error(exc):
                raise
            last_transient = exc
            logger.warning(
                f"Transient network error for {label} "
                f"(attempt {attempt + 1}/{len(attempts)}): {exc}"
            )
    assert last_transient is not None
    raise last_transient


class NetworkCircuit:
    """Pause once after N consecutive record-level transient failures."""

    def __init__(
        self,
        threshold: int = CIRCUIT_BREAKER_THRESHOLD,
        pause_seconds: float = CIRCUIT_BREAKER_PAUSE_SECONDS,
        sleep_fn=time.sleep,
    ):
        self.threshold = threshold
        self.pause_seconds = pause_seconds
        self.sleep_fn = sleep_fn
        self.consecutive = 0

    def note_success(self) -> None:
        self.consecutive = 0

    def note_transient_failure(self) -> None:
        self.consecutive += 1
        if self.consecutive < self.threshold:
            return
        announce_network_wait(
            self.pause_seconds,
            f"after {self.consecutive} consecutive failures",
        )
        self.sleep_fn(self.pause_seconds)
        self.consecutive = 0


def text_field(fields: dict, candidates: list[str], default: str = "") -> str:
    """Return the first populated Airtable field from possible field names."""
    for key in candidates:
        value = fields.get(key)
        if value is None:
            continue
        value = str(value).strip()
        if value:
            return value
    return default


def list_field(fields: dict, key: str) -> list:
    value = fields.get(key, [])
    if isinstance(value, list):
        return value
    return [value] if value else []


def table_count(supabase, table_name: str) -> int:
    def _count():
        result = supabase.table(table_name).select("id", count="exact").limit(1).execute()
        return result.count if result.count is not None else len(result.data or [])

    return retry_on_transient(_count, label=f"{table_name} count")


def existing_embedding_row_id(
    supabase,
    table_name: str,
    lookup_column: str,
    lookup_value: str,
) -> str:
    """Return an existing embedding row id, or '' if none."""

    def _lookup():
        existing = supabase.table(table_name).select("id").eq(
            lookup_column, lookup_value
        ).execute()
        rows = existing.data or []
        if not rows:
            return ""
        return rows[0].get("id") or ""

    return retry_on_transient(
        _lookup, label=f"existing {table_name} row for {lookup_value}"
    )


def store_single_embedding_row(
    supabase,
    table_name: str,
    lookup_column: str,
    lookup_value: str,
    payload: dict,
    preferred_id: str = "",
) -> tuple[str, str, int]:
    """
    Store exactly one embedding row for a source Airtable record.

    Supabase upsert only deduplicates when the database has a matching unique
    constraint. This project table currently allows duplicates, so we repair by
    updating one row and deleting stale duplicates for the same Airtable ID.
    """
    existing = supabase.table(table_name).select("id").eq(
        lookup_column, lookup_value
    ).execute()
    rows = existing.data or []

    duplicate_count = 0
    if rows:
        keep = next(
            (row for row in rows if preferred_id and row.get("id") == preferred_id),
            rows[0],
        )
        keep_id = keep["id"]
        supabase.table(table_name).update(payload).eq("id", keep_id).execute()

        duplicate_ids = [
            row["id"] for row in rows
            if row.get("id") and row.get("id") != keep_id
        ]
        duplicate_count = len(duplicate_ids)
        for duplicate_id in duplicate_ids:
            supabase.table(table_name).delete().eq("id", duplicate_id).execute()
        action = "updated"
    else:
        inserted = supabase.table(table_name).insert(payload).execute()
        keep_id = inserted.data[0]["id"] if inserted.data else ""
        action = "inserted"

    verification = supabase.table(table_name).select("id").eq(
        lookup_column, lookup_value
    ).execute()
    verified_rows = verification.data or []

    if len(verified_rows) != 1:
        raise RuntimeError(
            f"{table_name} verification failed for {lookup_value}: "
            f"expected 1 row, found {len(verified_rows)}"
        )

    verified_id = verified_rows[0]["id"]
    if keep_id and verified_id != keep_id:
        raise RuntimeError(
            f"{table_name} verification returned unexpected row for {lookup_value}"
        )

    return verified_id, action, duplicate_count

# ── CLIENTS ───────────────────────────────────────────────────────────────────
def get_supabase():
    from supabase import create_client

    return create_client(
        os.getenv("SUPABASE_URL"),
        os.getenv("SUPABASE_SERVICE_KEY"),
    )


def get_clients():
    """Supabase only. Kept so older call sites still receive a client."""
    return get_supabase()


def consultant_embed_records(supabase) -> list[dict]:
    """One embed record per stored id, including legacy embedding ids."""
    rows = []
    start = 0
    while True:
        res = (
            supabase.table("consultants")
            .select(
                "id,full_name,role_title,cv_text,thematic_expertise,"
                "geographic_experience,languages,tools,seniority_level,"
                "years_experience,day_rate_usd,availability_status,based_in,"
                "key_skills,legacy_ids,embedding_id"
            )
            .range(start, start + 999)
            .execute()
        )
        batch = res.data or []
        rows.extend(batch)
        if len(batch) < 1000:
            break
        start += 1000
    records = []
    for row in rows:
        ids = [str(item) for item in (row.get("legacy_ids") or []) if item]
        if row.get("id") and str(row["id"]) not in ids:
            ids.append(str(row["id"]))
        fields = {
            "full_name": row.get("full_name") or "",
            "role_title": row.get("role_title") or "",
            "cv_text": row.get("cv_text") or "",
            "thematic_expertise": row.get("thematic_expertise") or [],
            "geographic_experience": row.get("geographic_experience") or [],
            "languages": row.get("languages") or [],
            "tools": row.get("tools") or [],
            "seniority_level": row.get("seniority_level") or "",
            "years_experience": row.get("years_experience") or 0,
            "day_rate_usd": row.get("day_rate_usd") or 0,
            "availability_status": row.get("availability_status") or "",
            "based_in": row.get("based_in") or "",
            "key_skills": row.get("key_skills") or "",
            "embedding_id": row.get("embedding_id") or "",
        }
        for record_id in ids:
            records.append({"id": record_id, "fields": fields})
    return records


def proposal_files() -> list[Path]:
    folder = Path("data/proposals")
    if not folder.is_dir():
        return []
    return sorted(
        path for path in folder.iterdir()
        if path.is_file() and path.suffix.lower() in {".pdf", ".docx", ".txt", ".md"}
        and "PUT_PROPOSAL" not in path.name.upper()
    )


# ── EMBEDDING ─────────────────────────────────────────────────────────────────
def get_embedding_openai(text: str) -> list[float]:
    """
    Get embedding using OpenAI API.
    This produces 1536-dimension vectors compatible with our Supabase schema.
    
    NOTE: Requires OPENAI_API_KEY in .env
    Only costs ~$0.0001 per embedding — extremely cheap.
    """
    from config import get_openai_api_key

    api_key = get_openai_api_key()
    if not api_key:
        raise ValueError("OPENAI_API_KEY not set in .env")

    response = httpx.post(
        "https://api.openai.com/v1/embeddings",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json={
            "input": text[:8000],  # Max input length
            "model": EMBEDDING_MODEL,  # 1536 dims, cheap
        },
        timeout=30,
    )
    response.raise_for_status()
    return response.json()["data"][0]["embedding"]


# ── EMBED AND STORE CVS ───────────────────────────────────────────────────────
def embed_consultant_cv(
    supabase,
    consultant_table,
    consultant_record: dict,
    use_openai: bool = True,
    force: bool = False,
    circuit: NetworkCircuit | None = None,
) -> str:
    """Embed a consultant CV and store one verified row in Supabase pgvector."""

    fields = consultant_record.get("fields", {})
    airtable_id = consultant_record["id"]
    name = text_field(
        fields,
        ["full_name", "Full Name", "Name", "consultant_name"],
        default=f"Airtable record {airtable_id}",
    )
    cv_text = text_field(fields, ["cv_text", "CV Text", "cv", "resume_text"])

    if not fields:
        logger.warning(f"Blank Airtable consultant record {airtable_id}. Skipping.")
        return "skipped"

    if len(cv_text) < MIN_CV_TEXT_CHARS:
        logger.warning(
            f"No usable cv_text for {name} ({airtable_id}). Skipping embedding."
        )
        return "skipped"

    if not use_openai:
        logger.error("OPENAI_API_KEY is required to create CV vector embeddings.")
        return "failed"

    if not force:
        try:
            existing_id = existing_embedding_row_id(
                supabase,
                "cv_embeddings",
                "airtable_consultant_id",
                airtable_id,
            )
        except Exception as e:
            logger.error(f"Embedding failed for {name}: {e}")
            if circuit is not None and is_transient_network_error(e):
                circuit.note_transient_failure()
            return "failed"
        if existing_id:
            logger.info(f"Already embedded, skipping {name} ({existing_id})")
            if circuit is not None:
                circuit.note_success()
            return "skipped"

    metadata = {
        "thematic_expertise": list_field(fields, "thematic_expertise"),
        "geographic_experience": list_field(fields, "geographic_experience"),
        "languages": list_field(fields, "languages"),
        "tools": list_field(fields, "tools"),
        "seniority_level": fields.get("seniority_level", ""),
        "years_experience": fields.get("years_experience", 0),
        "day_rate_usd": fields.get("day_rate_usd", 0),
        "availability_status": fields.get("availability_status") or "",
        "based_in": fields.get("based_in", ""),
        "key_skills": fields.get("key_skills", ""),
    }

    def _embed_and_store():
        embedding = get_embedding_openai(cv_text)
        return store_single_embedding_row(
            supabase=supabase,
            table_name="cv_embeddings",
            lookup_column="airtable_consultant_id",
            lookup_value=airtable_id,
            preferred_id=fields.get("embedding_id", ""),
            payload={
                "airtable_consultant_id": airtable_id,
                "consultant_name": name,
                "role_title": fields.get("role_title", ""),
                "content_chunk": cv_text[:2000],
                "embedding": embedding,
                "metadata": metadata,
            },
        )

    try:
        embedding_id, action, duplicate_count = retry_on_transient(
            _embed_and_store, label=name
        )

        try:
            supabase.table("consultants").update(
                {"embedding_id": embedding_id}
            ).contains("legacy_ids", [airtable_id]).execute()
        except Exception as e:
            logger.warning(
                f"Supabase embedding stored for {name}, but consultants."
                f"embedding_id update failed: {e}"
            )
        _ = consultant_table

        if duplicate_count:
            logger.warning(
                f"Repaired {duplicate_count} duplicate CV embedding row(s) for {name}."
            )
        logger.success(f"CV embedding {action}: {name} ({embedding_id})")
        if circuit is not None:
            circuit.note_success()
        return "embedded"

    except Exception as e:
        logger.error(f"Embedding failed for {name}: {e}")
        if circuit is not None and is_transient_network_error(e):
            circuit.note_transient_failure()
        return "failed"


def embed_proposal(
    supabase,
    proposal_record: dict,
    use_openai: bool = True,
    force: bool = False,
    circuit: NetworkCircuit | None = None,
) -> str:
    """Embed a past proposal for style reference search."""

    fields = proposal_record.get("fields", {})
    airtable_id = proposal_record["id"]
    title = text_field(
        fields,
        ["project_title", "Project Title", "Name", "title"],
        default=f"Airtable record {airtable_id}",
    )
    proposal_text = text_field(fields, ["proposal_text", "Proposal Text", "text"])

    if not fields:
        logger.warning(f"Blank Airtable proposal record {airtable_id}. Skipping.")
        return "skipped"

    if len(proposal_text) < MIN_PROPOSAL_TEXT_CHARS:
        logger.warning(f"No usable proposal text for {title}. Skipping.")
        return "skipped"

    if not use_openai:
        logger.error("OPENAI_API_KEY is required to create proposal embeddings.")
        return "failed"

    if not force:
        try:
            existing_id = existing_embedding_row_id(
                supabase,
                "proposal_embeddings",
                "airtable_proposal_id",
                airtable_id,
            )
        except Exception as e:
            logger.error(f"Proposal embedding failed for {title}: {e}")
            if circuit is not None and is_transient_network_error(e):
                circuit.note_transient_failure()
            return "failed"
        if existing_id:
            logger.info(f"Already embedded, skipping {title} ({existing_id})")
            if circuit is not None:
                circuit.note_success()
            return "skipped"

    metadata = {
        "client": fields.get("client", ""),
        "year": fields.get("year", 0),
        "won": fields.get("won", False),
        "thematic_areas": list_field(fields, "thematic_areas"),
        "location": list_field(fields, "location"),
    }

    def _embed_and_store():
        embedding = get_embedding_openai(proposal_text)
        return store_single_embedding_row(
            supabase=supabase,
            table_name="proposal_embeddings",
            lookup_column="airtable_proposal_id",
            lookup_value=airtable_id,
            payload={
                "airtable_proposal_id": airtable_id,
                "project_title": title,
                "won": fields.get("won", False),
                "content_chunk": proposal_text[:2000],
                "embedding": embedding,
                "metadata": metadata,
            },
        )

    try:
        embedding_id, action, duplicate_count = retry_on_transient(
            _embed_and_store, label=title
        )

        if duplicate_count:
            logger.warning(
                f"Repaired {duplicate_count} duplicate proposal embedding row(s) "
                f"for {title}."
            )
        logger.success(f"Proposal embedding {action}: {title} ({embedding_id})")
        if circuit is not None:
            circuit.note_success()
        return "embedded"

    except Exception as e:
        logger.error(f"Proposal embedding failed for {title}: {e}")
        if circuit is not None and is_transient_network_error(e):
            circuit.note_transient_failure()
        return "failed"


def read_docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    sections = []
    for para in doc.paragraphs:
        if para.text.strip():
            sections.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                sections.append(" | ".join(cells))
    return "\n".join(sections)


def embed_proposal_from_file(
    supabase,
    path: Path,
    *,
    won: bool,
    title: str,
    client: str,
    year: int,
) -> str:
    """Embed a local past-proposal file. Does not call Airtable."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        proposal_text = read_docx_text(path)
    elif suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            proposal_text = "\n\n".join(
                (page.extract_text() or "") for page in pdf.pages
            )
    else:
        proposal_text = path.read_text(encoding="utf-8", errors="ignore")

    if len(proposal_text) < MIN_PROPOSAL_TEXT_CHARS:
        logger.warning(f"No usable proposal text in {path.name}. Skipping.")
        return "skipped"

    lookup = f"local:{path.name}"
    title = title or path.stem
    metadata = {
        "client": client,
        "year": year,
        "won": won,
        "source_file": path.name,
    }

    def _embed_and_store():
        embedding = get_embedding_openai(proposal_text)
        return store_single_embedding_row(
            supabase=supabase,
            table_name="proposal_embeddings",
            lookup_column="airtable_proposal_id",
            lookup_value=lookup,
            payload={
                "airtable_proposal_id": lookup,
                "project_title": title,
                "won": won,
                "content_chunk": proposal_text[:2000],
                "embedding": embedding,
                "metadata": metadata,
            },
        )

    embedding_id, action, duplicate_count = retry_on_transient(
        _embed_and_store, label=title
    )
    if duplicate_count:
        logger.warning(
            f"Repaired {duplicate_count} duplicate proposal embedding row(s) "
            f"for {title}."
        )
    logger.success(f"Proposal embedding {action}: {title} ({embedding_id})")
    console.print(
        f"  [green]{action}[/green] {title} "
        f"({table_count(supabase, 'proposal_embeddings')} proposal_embeddings rows)"
    )
    return "embedded"


# ── MAIN ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        description="Embed CVs and past proposals into Supabase."
    )
    parser.add_argument(
        "--file",
        metavar="NAME",
        help="Embed one file from data/proposals/.",
    )
    parser.add_argument(
        "--won",
        action="store_true",
        help="Mark the --file proposal as won.",
    )
    parser.add_argument(
        "--title",
        default="",
        help="project_title for --file (defaults to filename).",
    )
    parser.add_argument(
        "--client",
        default="",
        help="client metadata for --file.",
    )
    parser.add_argument(
        "--year",
        type=int,
        default=0,
        help="year metadata for --file.",
    )
    parser.add_argument(
        "--proposals-only",
        action="store_true",
        help="Skip CVs; embed past proposals from data/proposals/ only.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Re-embed even when a cv_embeddings / proposal_embeddings row already exists.",
    )
    args = parser.parse_args()

    console.print("\n[bold blue]Cortech CV Embedder[/bold blue]")

    use_openai = bool(os.getenv("OPENAI_API_KEY"))
    if not use_openai:
        console.print(
            "  OPENAI_API_KEY is required — CV matching depends on "
            "1536-dimension vector embeddings"
        )
        sys.exit(1)
    console.print("  OpenAI API key found — using vector embeddings")

    if not os.getenv("SUPABASE_URL") or not os.getenv("SUPABASE_SERVICE_KEY"):
        console.print(
            "\n[yellow]Supabase not configured.[/yellow]\n"
            "   Add SUPABASE_URL and SUPABASE_SERVICE_KEY to .env\n"
        )
        sys.exit(1)

    if args.file:
        proposals_dir = Path("data/proposals")
        wanted = args.file.lower()
        matches = [
            p for p in proposals_dir.iterdir()
            if p.is_file() and (
                p.name.lower() == wanted or wanted in p.name.lower()
            )
        ]
        if not matches:
            console.print(f"[red]No file matching '{args.file}' in {proposals_dir}[/red]")
            sys.exit(1)
        path = matches[0]
        console.print(f"  Embedding from disk: [cyan]{path.name}[/cyan]")
        try:
            status = embed_proposal_from_file(
                get_supabase(),
                path,
                won=args.won,
                title=args.title,
                client=args.client,
                year=args.year or datetime.now().year,
            )
        except Exception as e:
            console.print(f"[red]Embedding failed: {e}[/red]")
            sys.exit(1)
        if status != "embedded":
            sys.exit(1)
        console.print("\n[bold green]Embedding complete![/bold green]")
        return

    try:
        supabase = get_clients()
    except Exception as e:
        console.print(f"[red]Connection failed: {e}[/red]")
        sys.exit(1)

    cv_counts = {"embedded": 0, "skipped": 0, "failed": 0}
    prop_counts = {"embedded": 0, "skipped": 0, "failed": 0}
    circuit = NetworkCircuit()

    if not args.proposals_only:
        console.print("[bold]Loading consultants from Supabase...[/bold]")
        try:
            consultants = consultant_embed_records(supabase)
        except Exception as e:
            console.print(f"[red]Could not list consultants: {e}[/red]")
            sys.exit(1)
        console.print(f"Found [green]{len(consultants)}[/green] consultant id(s)\n")

        for record in track(consultants, description="Embedding CVs..."):
            status = embed_consultant_cv(
                supabase,
                None,
                record,
                use_openai,
                force=args.force,
                circuit=circuit,
            )
            cv_counts[status] += 1
            if status == "embedded":
                time.sleep(0.3)

    console.print("\n[bold]Loading past proposals from data/proposals/...[/bold]")
    proposals = proposal_files()
    console.print(f"Found [green]{len(proposals)}[/green] proposal file(s)\n")

    for path in track(proposals, description="Embedding proposals..."):
        lookup = f"local:{path.name}"
        if not args.force:
            try:
                existing_id = existing_embedding_row_id(
                    supabase,
                    "proposal_embeddings",
                    "airtable_proposal_id",
                    lookup,
                )
            except Exception as e:
                console.print(f"[red]Lookup failed for {path.name}: {e}[/red]")
                prop_counts["failed"] += 1
                continue
            if existing_id:
                prop_counts["skipped"] += 1
                continue
        try:
            status = embed_proposal_from_file(
                supabase,
                path,
                won=False,
                title=path.stem,
                client="",
                year=datetime.now().year,
            )
        except Exception as e:
            console.print(f"[red]Embedding failed for {path.name}: {e}[/red]")
            status = "failed"
        prop_counts[status] = prop_counts.get(status, 0) + 1
        if status == "embedded":
            time.sleep(0.3)

    console.print(
        f"\n  Proposals: [green]{prop_counts['embedded']} embedded[/green] | "
        f"[yellow]{prop_counts['skipped']} skipped[/yellow] | "
        f"[red]{prop_counts['failed']} failed[/red]"
    )
    console.print(
        f"  Supabase proposal_embeddings rows: "
        f"[cyan]{table_count(supabase, 'proposal_embeddings')}[/cyan]"
    )

    if not args.proposals_only:
        console.print(
            f"\n  CVs: [green]{cv_counts['embedded']} embedded[/green] | "
            f"[yellow]{cv_counts['skipped']} skipped[/yellow] | "
            f"[red]{cv_counts['failed']} failed[/red]"
        )
        console.print(
            f"  Supabase cv_embeddings rows: "
            f"[cyan]{table_count(supabase, 'cv_embeddings')}[/cyan]"
        )

    failed_parts = []
    if cv_counts["failed"]:
        failed_parts.append(f"{cv_counts['failed']} CV(s)")
    if prop_counts["failed"]:
        failed_parts.append(f"{prop_counts['failed']} proposal(s)")
    if failed_parts:
        console.print(
            f"\n[bold red]{' and '.join(failed_parts)} failed after retries.[/bold red]\n"
            "  Re-run: [cyan]python embed_cvs.py[/cyan]\n"
            "  Already-embedded CVs will be skipped; only remaining failures retry.\n"
            "  Use [cyan]python embed_cvs.py --force[/cyan] to re-embed everything."
        )
        sys.exit(1)

    console.print("\n[bold green]Embedding complete![/bold green]")
    console.print("\n[bold]Next step:[/bold]")
    console.print("  Run: [cyan]python main.py --once[/cyan]")


if __name__ == "__main__":
    main()